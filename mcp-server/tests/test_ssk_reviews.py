import json
from unittest.mock import patch

from docx import Document

from tools import ssk_common
from tools.ssk import ssk_import_catalog, ssk_status
from tools.ssk_evidence import ssk_link_evidence, ssk_verify_pointers
from tools.ssk_reviews import (
    ssk_alerts,
    ssk_due,
    ssk_record_review,
    ssk_review_history,
    ssk_review_notifications,
)


def _seed_control_via_import(tmp_path, control_id: str = "06-3") -> None:
    category, suffix = control_id.split("-", 1)
    docx_path = tmp_path / f"{control_id}.docx"
    doc = Document()
    doc.add_heading("MCNA Cyber Security Standards", level=1)
    doc.add_heading(
        f"{control_id}: Imported control {suffix}"
        f"SSK Group {category}: Test StandardsNOF MCNA...",
        level=1,
    )
    for key, val in [("Effective Date", "2026-05-15"), ("Review Date", "2027-05-15"),
                     ("Approver", "CFO / Executive Sponsor")]:
        doc.add_paragraph(key)
        doc.add_paragraph(val)
    doc.add_heading("1. Purpose", level=3)
    doc.add_paragraph("Overview text.")
    doc.add_heading("4. Standards", level=3)
    doc.add_paragraph("The obligations below apply to MCNA.")
    doc.add_heading("G", level=4)
    doc.add_paragraph(f"{control_id}.1Action one.")
    doc.add_heading("6. Review & Compliance", level=3)
    doc.add_heading("Cadence", level=4)
    doc.add_paragraph("Status text.")
    doc.add_heading("Reviewer", level=4)
    doc.add_paragraph("IT-MIS Director.")
    doc.add_heading("Audit evidence", level=4)
    p = doc.add_paragraph("Evidence item.")
    p.style = doc.styles["List Paragraph"]
    doc.add_heading("7. Risks of Non-compliance", level=3)
    doc.add_paragraph("Risk text.")
    doc.save(docx_path)
    ssk_import_catalog(
        docx_path=str(docx_path),
        version=f"seed-{category}",
        dry_run=False,
        intermediate_dir=str(tmp_path / "intermediate"),
    )


def test_ssk_record_review_rejects_mixed_review_scope(db, tmp_path):
    _seed_control_via_import(tmp_path, control_id="06-3")

    result = json.loads(
        ssk_record_review(
            control_id="06-3",
            control_family="06",
            reviewer="tester@example.com",
            evidence_ids=[],
            outcome="ok",
        )
    )

    assert result["error_type"] == "ReviewScopeError"


def test_ssk_record_review_without_evidence_sets_no_evidence_quality_flag(db, tmp_path, monkeypatch):
    _seed_control_via_import(tmp_path, control_id="06-3")
    reviewed_at = "2026-04-22T12:00:00+00:00"
    monkeypatch.setattr(ssk_common, "utc_now", lambda: reviewed_at)

    result = json.loads(
        ssk_record_review(
            control_id="06-3",
            reviewer="tester@example.com",
            evidence_ids=[],
            scope_summary="Quarterly review",
            outcome="ok",
            findings_summary="No linked evidence yet.",
        )
    )

    assert result["control_id"] == "06-3"
    assert result["quality_flag"] == "no_evidence"
    assert result["reviewed_at"] == reviewed_at


def test_ssk_record_review_with_current_evidence_updates_status_to_regularly_reviewed(
    db, tmp_path, monkeypatch
):
    _seed_control_via_import(tmp_path, control_id="06-3")
    produced_at = "2026-04-10T12:00:00+00:00"
    reviewed_at = "2026-04-22T12:00:00+00:00"

    monkeypatch.setattr(ssk_common, "utc_now", lambda: produced_at)
    evidence_one = json.loads(
        ssk_link_evidence(
            control_id="06-3",
            evidence_type="review_minutes",
            source_kind="local_file",
            source_pointer="C:/evidence/review-minutes.docx",
            validity_window_days=30,
            title="Review minutes",
        )
    )
    evidence_two = json.loads(
        ssk_link_evidence(
            control_id="06-3",
            evidence_type="policy_link",
            source_kind="local_file",
            source_pointer="C:/evidence/review-policy.docx",
            validity_window_days=30,
            title="Review policy",
        )
    )

    monkeypatch.setattr(ssk_common, "utc_now", lambda: reviewed_at)
    review = json.loads(
        ssk_record_review(
            control_id="06-3",
            reviewer="tester@example.com",
            evidence_ids=[evidence_one["evidence_id"], evidence_two["evidence_id"]],
            scope_summary="Quarterly review",
            outcome="ok",
            findings_summary="Evidence current and validated for review.",
        )
    )
    status = json.loads(ssk_status("06-3"))

    assert review["quality_flag"] == "ok"
    assert status["current_maturity"] == "regularly_reviewed"
    assert status["last_reviewed_at"] == reviewed_at


def test_ssk_verify_pointers_invalidates_review_status_and_alerts(db, tmp_path, monkeypatch):
    _seed_control_via_import(tmp_path, control_id="06-3")
    evidence_path_one = tmp_path / "review-minutes.docx"
    evidence_path_two = tmp_path / "review-policy.docx"
    evidence_path_one.write_text("minutes", encoding="utf-8")
    evidence_path_two.write_text("policy", encoding="utf-8")

    produced_at = "2026-04-10T12:00:00+00:00"
    reviewed_at = "2026-04-22T12:00:00+00:00"
    invalidated_at = "2026-04-23T09:00:00+00:00"

    monkeypatch.setattr(ssk_common, "utc_now", lambda: produced_at)
    evidence_one = json.loads(
        ssk_link_evidence(
            control_id="06-3",
            evidence_type="review_minutes",
            source_kind="local_file",
            source_pointer=str(evidence_path_one),
            validity_window_days=30,
            title="Review minutes",
        )
    )
    evidence_two = json.loads(
        ssk_link_evidence(
            control_id="06-3",
            evidence_type="policy_link",
            source_kind="local_file",
            source_pointer=str(evidence_path_two),
            validity_window_days=30,
            title="Review policy",
        )
    )

    monkeypatch.setattr(ssk_common, "utc_now", lambda: reviewed_at)
    review = json.loads(
        ssk_record_review(
            control_id="06-3",
            reviewer="tester@example.com",
            evidence_ids=[evidence_one["evidence_id"], evidence_two["evidence_id"]],
            scope_summary="Quarterly review",
            outcome="ok",
            findings_summary="Evidence current and validated for review.",
        )
    )
    assert review["quality_flag"] == "ok"
    assert json.loads(ssk_alerts()) == []
    assert json.loads(ssk_status("06-3"))["current_maturity"] == "regularly_reviewed"

    evidence_path_one.unlink()
    monkeypatch.setattr(ssk_common, "utc_now", lambda: invalidated_at)
    verify_result = json.loads(ssk_verify_pointers(control_id="06-3"))
    status = json.loads(ssk_status("06-3"))
    alerts = json.loads(ssk_alerts())
    history = json.loads(ssk_review_history("06-3"))

    assert verify_result["broken"] == 1
    assert status["current_maturity"] == "not_regularly_reviewed"
    assert alerts[0]["review_id"] == review["review_id"]
    assert alerts[0]["quality_flag"] == "stale_evidence"
    assert history[0]["review_id"] == review["review_id"]
    assert history[0]["quality_flag"] == "stale_evidence"


def test_ssk_refresh_downgrades_review_after_evidence_expires(db, tmp_path, monkeypatch):
    _seed_control_via_import(tmp_path, control_id="06-3")
    evidence_path_one = tmp_path / "review-minutes.docx"
    evidence_path_two = tmp_path / "review-policy.docx"
    evidence_path_one.write_text("minutes", encoding="utf-8")
    evidence_path_two.write_text("policy", encoding="utf-8")

    produced_at = "2026-04-10T12:00:00+00:00"
    reviewed_at = "2026-04-12T12:00:00+00:00"
    refreshed_at = "2026-04-22T09:00:00+00:00"

    monkeypatch.setattr(ssk_common, "utc_now", lambda: produced_at)
    evidence_one = json.loads(
        ssk_link_evidence(
            control_id="06-3",
            evidence_type="review_minutes",
            source_kind="local_file",
            source_pointer=str(evidence_path_one),
            validity_window_days=5,
            title="Review minutes",
        )
    )
    evidence_two = json.loads(
        ssk_link_evidence(
            control_id="06-3",
            evidence_type="policy_link",
            source_kind="local_file",
            source_pointer=str(evidence_path_two),
            validity_window_days=5,
            title="Review policy",
        )
    )

    monkeypatch.setattr(ssk_common, "utc_now", lambda: reviewed_at)
    review = json.loads(
        ssk_record_review(
            control_id="06-3",
            reviewer="tester@example.com",
            evidence_ids=[evidence_one["evidence_id"], evidence_two["evidence_id"]],
            scope_summary="Quarterly review",
            outcome="ok",
            findings_summary="Evidence was current at review time.",
        )
    )
    assert review["quality_flag"] == "ok"
    assert json.loads(ssk_status("06-3"))["current_maturity"] == "regularly_reviewed"
    assert json.loads(ssk_alerts()) == []

    monkeypatch.setattr(ssk_common, "utc_now", lambda: refreshed_at)
    verify_result = json.loads(ssk_verify_pointers(control_id="06-3"))
    status = json.loads(ssk_status("06-3"))
    alerts = json.loads(ssk_alerts())
    history = json.loads(ssk_review_history("06-3"))

    assert verify_result["resolved"] == 2
    assert verify_result["broken"] == 0
    assert status["current_maturity"] == "not_regularly_reviewed"
    assert alerts[0]["review_id"] == review["review_id"]
    assert alerts[0]["quality_flag"] == "stale_evidence"
    assert history[0]["review_id"] == review["review_id"]
    assert history[0]["quality_flag"] == "stale_evidence"


def test_ssk_review_tools_canonicalize_control_aliases(db, tmp_path, monkeypatch):
    _seed_control_via_import(tmp_path, control_id="08-1")
    reviewed_at = "2026-04-22T12:00:00+00:00"
    monkeypatch.setattr(ssk_common, "utc_now", lambda: reviewed_at)

    review = json.loads(
        ssk_record_review(
            control_id="CA-COV-01",
            reviewer="tester@example.com",
            evidence_ids=[],
            scope_summary="Alias-based review",
            outcome="ok",
            findings_summary="Alias should resolve to canonical control.",
        )
    )
    history = json.loads(ssk_review_history("CA-COV-01"))

    assert review["control_id"] == "08-1"
    assert history[0]["control_id"] == "08-1"
    assert history[0]["review_id"] == review["review_id"]


def test_ssk_due_returns_never_reviewed_by_default(db, tmp_path):
    _seed_control_via_import(tmp_path, control_id="06-3")

    result = json.loads(ssk_due(days_ahead=30))

    assert len(result) == 1
    assert result[0]["control_id"] == "06-3"
    assert result[0]["review_status"] == "never_reviewed"
    assert result[0]["next_review_due"] is None


def test_ssk_due_excludes_never_reviewed_when_flag_false(db, tmp_path):
    _seed_control_via_import(tmp_path, control_id="06-3")

    result = json.loads(ssk_due(days_ahead=30, include_never_reviewed=False))

    assert result == []


def test_ssk_due_overdue_after_review_recorded(db, tmp_path, monkeypatch):
    _seed_control_via_import(tmp_path, control_id="06-3")
    # Review recorded in the past; next_review_due = 2026-01-01 + 90d = 2026-04-01
    monkeypatch.setattr(ssk_common, "utc_now", lambda: "2026-01-01T12:00:00+00:00")
    ssk_record_review(
        control_id="06-3",
        reviewer="tester@example.com",
        evidence_ids=[],
        scope_summary="Initial review",
        outcome="ok",
    )
    # Advance clock past next_review_due so the control shows as overdue
    monkeypatch.setattr(ssk_common, "utc_now", lambda: "2026-05-01T12:00:00+00:00")

    result = json.loads(ssk_due(days_ahead=30))

    assert len(result) == 1
    assert result[0]["control_id"] == "06-3"
    assert result[0]["review_status"] == "overdue"
    assert result[0]["next_review_due"] is not None


def test_ssk_due_not_returned_when_review_is_current(db, tmp_path, monkeypatch):
    _seed_control_via_import(tmp_path, control_id="06-3")
    # Review recorded today — next_review_due is 90 days out, well past days_ahead=30
    monkeypatch.setattr(ssk_common, "utc_now", lambda: "2026-04-26T12:00:00+00:00")
    ssk_record_review(
        control_id="06-3",
        reviewer="tester@example.com",
        evidence_ids=[],
        scope_summary="Current review",
        outcome="ok",
    )

    result = json.loads(ssk_due(days_ahead=30, include_never_reviewed=False))

    assert result == []


def test_ssk_review_notifications_returns_report_without_mail(db, tmp_path, monkeypatch):
    _seed_control_via_import(tmp_path, control_id="06-3")
    monkeypatch.setattr(ssk_common, "utc_now", lambda: "2026-04-27T12:00:00+00:00")

    result = json.loads(ssk_review_notifications(days_ahead=30))

    assert result["generated_at"] == "2026-04-27T12:00:00+00:00"
    assert result["counts"] == {
        "total": 1,
        "overdue": 0,
        "due_soon": 0,
        "never_reviewed": 1,
    }
    assert result["controls"][0]["control_id"] == "06-3"
    assert result["mail"] is None
    assert "06-3" in result["body"]
    assert "never reviewed" in result["body"].lower()


def test_ssk_review_notifications_uses_mail_send_summary_for_dry_run_preview(
    db, tmp_path, monkeypatch
):
    _seed_control_via_import(tmp_path, control_id="06-3")
    monkeypatch.setattr(ssk_common, "utc_now", lambda: "2026-04-27T12:00:00+00:00")

    with patch(
        "tools.ssk_reviews.mail_send_summary",
        return_value=json.dumps({"dry_run": True, "sent": False, "subject": "preview"}),
    ) as mail_send:
        result = json.loads(
            ssk_review_notifications(
                days_ahead=30,
                to="ops@nofmetalcoatings.us",
                cc="audit@nofmetalcoatings.us",
            )
        )

    mail_send.assert_called_once()
    kwargs = mail_send.call_args.kwargs
    assert kwargs["to"] == "ops@nofmetalcoatings.us"
    assert kwargs["cc"] == "audit@nofmetalcoatings.us"
    assert kwargs["dry_run"] is True
    assert "MCNA Secure SketCH review queue" in kwargs["subject"]
    assert "06-3" in kwargs["body"]
    assert result["mail"]["dry_run"] is True


def test_ssk_review_notifications_registered_in_server():
    import server

    assert server.ssk_review_notifications is ssk_review_notifications
