import pytest
from docx import Document
from tools.ssk_parser import parse_catalog, CatalogSourceError


def _build_minimal_mcna_docx(tmp_path, controls=None):
    """Build a minimal MCNA consolidated-format .docx for parser tests."""
    if controls is None:
        controls = [
            {
                "control_id": "04-1",
                "title": "Sample HR control",
                "category_suffix": "Human Resources Standards",
                "effective_date": "2026-05-15",
                "review_date": "2027-05-15",
                "approver": "CFO / Executive Sponsor",
                "purpose": "Purpose text for 04-1.",
                "clause_groups": [
                    ("Governance", [("04-1.1", "Clause one."), ("04-1.2", "Clause two.")]),
                ],
                "cadence": "Annually.",
                "reviewer": "IT-MIS Director.",
                "audit_evidence": ["Evidence item one.", "Evidence item two."],
                "risks": "Risk text for 04-1.",
            }
        ]
    path = tmp_path / "minimal_mcna.docx"
    doc = Document()
    doc.add_heading("MCNA Cyber Security Standards", level=1)
    for c in controls:
        cid = c["control_id"]
        gnum = cid.split("-")[0]
        doc.add_heading(
            f"{cid}: {c['title']}SSK Group {gnum}: {c['category_suffix']}"
            "NOF MCNA IT and Information Systems Security Policy and Standards",
            level=1,
        )
        for key, val in [
            ("Document ID", f"IT-STD-SSK-{cid}"),
            ("Effective Date", c["effective_date"]),
            ("Review Date", c["review_date"]),
            ("Approver", c["approver"]),
        ]:
            doc.add_paragraph(key)
            doc.add_paragraph(val)
        doc.add_heading("1. Purpose", level=3)
        doc.add_paragraph(c["purpose"])
        doc.add_heading("4. Standards", level=3)
        doc.add_paragraph("The obligations below apply to MCNA.")
        current_group = None
        for group_name, clauses in c["clause_groups"]:
            if group_name != current_group:
                doc.add_heading(group_name, level=4)
                current_group = group_name
            for clause_id, clause_text in clauses:
                doc.add_paragraph(f"{clause_id}{clause_text}")
        doc.add_heading("6. Review & Compliance", level=3)
        doc.add_heading("Cadence", level=4)
        doc.add_paragraph(c["cadence"])
        doc.add_heading("Reviewer", level=4)
        doc.add_paragraph(c["reviewer"])
        doc.add_heading("Audit evidence", level=4)
        for item in c["audit_evidence"]:
            p = doc.add_paragraph(item)
            p.style = doc.styles["List Paragraph"]
        doc.add_heading("7. Risks of Non-compliance", level=3)
        doc.add_paragraph(c["risks"])
    doc.save(path)
    return path


def test_parser_identifies_control(tmp_path):
    path = _build_minimal_mcna_docx(tmp_path)
    result = parse_catalog(path, source_version="test-1")
    assert len(result["controls"]) == 1
    assert result["controls"][0]["control_id"] == "04-1"
    assert result["parse_failures"] == []


def test_parser_extracts_metadata(tmp_path):
    path = _build_minimal_mcna_docx(tmp_path)
    result = parse_catalog(path, source_version="test-1")
    c = result["controls"][0]
    assert c["title"] == "Sample HR control"
    assert c["effective_date"] == "2026-05-15"
    assert c["review_date"] == "2027-05-15"
    assert c["approver"] == "CFO / Executive Sponsor"
    assert c["category"] == "04"
    assert c["category_name"] is not None and "Human Resources" in c["category_name"]


def test_parser_extracts_overview(tmp_path):
    path = _build_minimal_mcna_docx(tmp_path)
    result = parse_catalog(path, source_version="test-1")
    assert result["controls"][0]["overview"] == "Purpose text for 04-1."


def test_parser_extracts_clauses(tmp_path):
    path = _build_minimal_mcna_docx(tmp_path)
    result = parse_catalog(path, source_version="test-1")
    clauses = result["controls"][0]["clauses"]
    assert len(clauses) == 2
    assert clauses[0]["clause_id"] == "04-1.1"
    assert clauses[0]["clause_text"] == "Clause one."
    assert clauses[0]["group_name"] == "Governance"
    assert clauses[0]["sequence"] == 0
    assert clauses[1]["clause_id"] == "04-1.2"
    assert clauses[1]["sequence"] == 1


def test_parser_extracts_review_section(tmp_path):
    path = _build_minimal_mcna_docx(tmp_path)
    result = parse_catalog(path, source_version="test-1")
    c = result["controls"][0]
    assert c["cadence"] == "Annually."
    assert c["reviewer"] == "IT-MIS Director."
    assert c["status_description"] == "Annually."  # cadence is aliased to status_description


def test_parser_extracts_audit_evidence(tmp_path):
    path = _build_minimal_mcna_docx(tmp_path)
    result = parse_catalog(path, source_version="test-1")
    items = result["controls"][0]["audit_evidence_items"]
    assert len(items) == 2
    assert items[0]["item_text"] == "Evidence item one."
    assert items[0]["sequence"] == 0
    assert items[1]["item_text"] == "Evidence item two."
    assert items[1]["sequence"] == 1


def test_parser_extracts_risks(tmp_path):
    path = _build_minimal_mcna_docx(tmp_path)
    result = parse_catalog(path, source_version="test-1")
    assert "Risk text for 04-1." in result["controls"][0]["insufficient_measures_risks"]


def test_parser_reports_failures_for_missing_sections(tmp_path):
    path = tmp_path / "bad.docx"
    doc = Document()
    doc.add_heading(
        "04-1: Missing Sections ControlSSK Group 04: HR StandardsNOF MCNA...", level=1
    )
    doc.add_heading("7. Risks of Non-compliance", level=3)
    doc.add_paragraph("Some risk.")
    doc.save(path)
    result = parse_catalog(path, source_version="test-fail")
    assert len(result["parse_failures"]) == 1
    failure = result["parse_failures"][0]
    assert failure["control_id"] == "04-1"
    assert "overview" in failure["reason"]
    assert "standards clauses" in failure["reason"]


def test_parser_handles_two_controls(tmp_path):
    path = _build_minimal_mcna_docx(
        tmp_path,
        controls=[
            {
                "control_id": "04-1",
                "title": "HR control",
                "category_suffix": "Human Resources Standards",
                "effective_date": "2026-05-15",
                "review_date": "2027-05-15",
                "approver": "CFO / Executive Sponsor",
                "purpose": "HR purpose.",
                "clause_groups": [("G", [("04-1.1", "HR clause.")])],
                "cadence": "Annually.",
                "reviewer": "IT-MIS Director.",
                "audit_evidence": ["HR evidence."],
                "risks": "HR risk.",
            },
            {
                "control_id": "06-3",
                "title": "Asset control",
                "category_suffix": "Asset Management Standards",
                "effective_date": "2026-05-15",
                "review_date": "2027-05-15",
                "approver": "CFO / Executive Sponsor",
                "purpose": "Asset purpose.",
                "clause_groups": [
                    ("G", [("06-3.1", "Asset clause one."), ("06-3.2", "Asset clause two.")])
                ],
                "cadence": "Annually.",
                "reviewer": "IT-MIS Director.",
                "audit_evidence": ["Asset evidence one.", "Asset evidence two."],
                "risks": "Asset risk.",
            },
        ],
    )
    result = parse_catalog(path, source_version="test-two")
    assert len(result["controls"]) == 2
    assert [c["control_id"] for c in result["controls"]] == ["04-1", "06-3"]
    assert result["controls"][1]["clauses"][0]["clause_id"] == "06-3.1"


def test_parser_skips_cover_heading(tmp_path):
    """The first Heading1 ('MCNA Cyber Security Standards') must not be parsed as a control."""
    path = _build_minimal_mcna_docx(tmp_path)
    result = parse_catalog(path, source_version="test-cover")
    ids = [c["control_id"] for c in result["controls"]]
    assert "MCNA" not in ids
    assert len(result["controls"]) == 1


def test_parser_missing_file_raises(tmp_path):
    with pytest.raises(CatalogSourceError):
        parse_catalog(tmp_path / "does-not-exist.docx", source_version="x")
