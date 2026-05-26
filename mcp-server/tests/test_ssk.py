import json
from pathlib import Path
from docx import Document
from tools.ssk import ssk_import_catalog, ssk_list_controls, ssk_get_control, ssk_status, ssk_status_all, ssk_gaps
from tools.ssk_evidence import ssk_link_evidence, ssk_list_evidence
from tools.ssk_reviews import ssk_record_review, ssk_review_history
from tools.ssk_binder import ssk_coverage, ssk_export_binder

def _build_minimal_docx(tmp_path):
    """Single control, 1 clause, 1 audit evidence item — used for basic import tests."""
    path = tmp_path / "minimal.docx"
    doc = Document()
    doc.add_heading("MCNA Cyber Security Standards", level=1)
    doc.add_heading(
        "04-1: Sample controlSSK Group 04: Human Resources StandardsNOF MCNA...", level=1
    )
    for key, val in [
        ("Effective Date", "2026-05-15"),
        ("Review Date", "2027-05-15"),
        ("Approver", "CFO / Executive Sponsor"),
    ]:
        doc.add_paragraph(key)
        doc.add_paragraph(val)
    doc.add_heading("1. Purpose", level=3)
    doc.add_paragraph("Overview text.")
    doc.add_heading("4. Standards", level=3)
    doc.add_paragraph("The obligations below apply to MCNA.")
    doc.add_heading("Governance", level=4)
    doc.add_paragraph("04-1.1Action one.")
    doc.add_heading("6. Review & Compliance", level=3)
    doc.add_heading("Cadence", level=4)
    doc.add_paragraph("Status text.")
    doc.add_heading("Reviewer", level=4)
    doc.add_paragraph("IT-MIS Director.")
    doc.add_heading("Audit evidence", level=4)
    p = doc.add_paragraph("Audit evidence item.")
    p.style = doc.styles["List Paragraph"]
    doc.add_heading("7. Risks of Non-compliance", level=3)
    doc.add_paragraph("Risk text.")
    doc.save(path)
    return path


def test_ssk_import_catalog_dry_run_writes_json_only(db, tmp_path):
    from db import get_connection
    docx_path = _build_minimal_docx(tmp_path)
    intermediate_dir = tmp_path / "intermediate"
    result_json = ssk_import_catalog(
        docx_path=str(docx_path),
        version="test-1",
        dry_run=True,
        intermediate_dir=str(intermediate_dir),
    )
    result = json.loads(result_json)
    assert result["dry_run"] is True
    assert result["parsed_controls"] == 1
    assert result["written_controls"] == 0

    with get_connection() as conn:
        rows = conn.execute("SELECT COUNT(*) AS c FROM ssk_controls").fetchone()
    assert rows["c"] == 0

    intermediate = Path(result["intermediate_json_path"])
    assert intermediate.exists()
    payload = json.loads(intermediate.read_text(encoding="utf-8"))
    assert payload["controls"][0]["control_id"] == "04-1"


def test_ssk_import_catalog_commit_writes_db(db, tmp_path):
    from db import get_connection
    docx_path = _build_minimal_docx(tmp_path)
    intermediate_dir = tmp_path / "intermediate"
    result_json = ssk_import_catalog(
        docx_path=str(docx_path),
        version="test-1",
        dry_run=False,
        intermediate_dir=str(intermediate_dir),
    )
    result = json.loads(result_json)
    assert result["dry_run"] is False
    assert result["written_controls"] == 1
    assert result["written_clauses"] == 1
    assert result["written_evidence_items"] == 1

    with get_connection() as conn:
        rows = conn.execute("SELECT COUNT(*) AS c FROM ssk_controls").fetchone()
    assert rows["c"] == 1


def _seed_two_controls(db, tmp_path):
    path = tmp_path / "two.docx"
    doc = Document()
    doc.add_heading("MCNA Cyber Security Standards", level=1)
    # Control 04-1
    doc.add_heading(
        "04-1: HR controlSSK Group 04: Human Resources StandardsNOF MCNA...", level=1
    )
    for key, val in [("Effective Date", "2026-05-15"), ("Review Date", "2027-05-15"),
                     ("Approver", "CFO / Executive Sponsor")]:
        doc.add_paragraph(key)
        doc.add_paragraph(val)
    doc.add_heading("1. Purpose", level=3)
    doc.add_paragraph("HR overview.")
    doc.add_heading("4. Standards", level=3)
    doc.add_paragraph("The obligations below apply to MCNA.")
    doc.add_heading("G", level=4)
    doc.add_paragraph("04-1.1HR action.")
    doc.add_heading("6. Review & Compliance", level=3)
    doc.add_heading("Cadence", level=4)
    doc.add_paragraph("HR status.")
    doc.add_heading("Reviewer", level=4)
    doc.add_paragraph("IT-MIS Director.")
    doc.add_heading("Audit evidence", level=4)
    p = doc.add_paragraph("HR evidence.")
    p.style = doc.styles["List Paragraph"]
    doc.add_heading("7. Risks of Non-compliance", level=3)
    doc.add_paragraph("HR risk.")
    # Control 06-3
    doc.add_heading(
        "06-3: Asset controlSSK Group 06: Asset Management StandardsNOF MCNA...", level=1
    )
    for key, val in [("Effective Date", "2026-05-15"), ("Review Date", "2027-05-15"),
                     ("Approver", "CFO / Executive Sponsor")]:
        doc.add_paragraph(key)
        doc.add_paragraph(val)
    doc.add_heading("1. Purpose", level=3)
    doc.add_paragraph("Asset overview.")
    doc.add_heading("4. Standards", level=3)
    doc.add_paragraph("The obligations below apply to MCNA.")
    doc.add_heading("G", level=4)
    doc.add_paragraph("06-3.1Asset clause one.")
    doc.add_paragraph("06-3.2Asset clause two.")
    doc.add_heading("6. Review & Compliance", level=3)
    doc.add_heading("Cadence", level=4)
    doc.add_paragraph("Asset status.")
    doc.add_heading("Reviewer", level=4)
    doc.add_paragraph("IT-MIS Director.")
    doc.add_heading("Audit evidence", level=4)
    p = doc.add_paragraph("Asset evidence.")
    p.style = doc.styles["List Paragraph"]
    doc.add_heading("7. Risks of Non-compliance", level=3)
    doc.add_paragraph("Asset risk.")
    doc.save(path)
    ssk_import_catalog(str(path), version="seed", dry_run=False,
                       intermediate_dir=str(tmp_path / "i"))


def test_ssk_list_controls_returns_all(db, tmp_path):
    _seed_two_controls(db, tmp_path)
    result = json.loads(ssk_list_controls())
    ids = [c["control_id"] for c in result]
    assert ids == ["04-1", "06-3"]


def test_ssk_list_controls_filters_by_category(db, tmp_path):
    _seed_two_controls(db, tmp_path)
    result = json.loads(ssk_list_controls(category="06"))
    assert len(result) == 1
    assert result[0]["control_id"] == "06-3"


def test_ssk_get_control_includes_clauses(db, tmp_path):
    _seed_two_controls(db, tmp_path)
    result = json.loads(ssk_get_control("06-3"))
    assert result["control_id"] == "06-3"
    assert result["title"] == "Asset control"
    assert len(result["clauses"]) == 2
    assert result["clauses"][0]["clause_text"] == "Asset clause one."
    assert result["clauses"][0]["clause_id"] == "06-3.1"
    assert len(result["audit_evidence_items"]) == 1
    assert result["audit_evidence_items"][0]["item_text"] == "Asset evidence."


def test_ssk_get_control_unknown_returns_error(db):
    result = json.loads(ssk_get_control("99-9"))
    assert "error" in result


def test_ssk_status_default_after_import(db, tmp_path):
    _seed_two_controls(db, tmp_path)
    result = json.loads(ssk_status("06-3"))
    assert result["control_id"] == "06-3"
    assert result["current_maturity"] == "not_regularly_reviewed"
    assert result["target_maturity"] == "Regularly Reviewed"
    assert result["title"] == "Asset control"
    assert result["last_reviewed_at"] is None


def test_ssk_status_unknown_returns_error(db):
    result = json.loads(ssk_status("99-9"))
    assert "error" in result


def test_ssk_status_all_returns_every_imported_control(db, tmp_path):
    _seed_two_controls(db, tmp_path)
    result = json.loads(ssk_status_all())
    assert len(result) == 2
    ids = [r["control_id"] for r in result]
    assert ids == ["04-1", "06-3"]


def test_ssk_status_all_below_target_filter(db, tmp_path):
    _seed_two_controls(db, tmp_path)
    # All controls start below target after import — filter should return both
    result = json.loads(ssk_status_all(below_target=True))
    assert len(result) == 2


def test_ssk_gaps_lists_controls_below_target(db, tmp_path):
    _seed_two_controls(db, tmp_path)
    result = json.loads(ssk_gaps())
    assert result["total_controls"] == 2
    assert result["at_target"] == 0
    assert result["below_target"] == 2
    assert len(result["gaps"]) == 2
    gap_ids = {g["control_id"] for g in result["gaps"]}
    assert gap_ids == {"04-1", "06-3"}
    for g in result["gaps"]:
        assert g["current_maturity"] == "not_regularly_reviewed"
        assert "title" in g


def test_phase2b_flow_import_link_review_and_export(db, tmp_path):
    """Integration: import → link evidence → record review → export binder → coverage."""
    _seed_two_controls(db, tmp_path)

    # Link evidence to 04-1
    ev = json.loads(ssk_link_evidence(
        control_id="04-1",
        source_pointer="file://test/hr_policy.pdf",
        source_kind="local_file",
        evidence_type="policy_link",
        title="HR policy doc",
    ))
    assert ev["evidence_id"] is not None
    ev_id = ev["evidence_id"]

    # Confirm evidence listed
    listed = json.loads(ssk_list_evidence("04-1"))
    assert len(listed) == 1
    assert listed[0]["source_pointer"] == "file://test/hr_policy.pdf"

    # Record a review on 04-1
    rev = json.loads(ssk_record_review(
        control_id="04-1",
        outcome="ok",
        reviewer="dlafferty",
        findings_summary="Looks good.",
    ))
    assert rev["review_id"] is not None
    assert rev["outcome"] == "ok"

    # Confirm review history
    hist = json.loads(ssk_review_history("04-1"))
    assert len(hist) == 1
    assert hist[0]["outcome"] == "ok"

    # Export binder for 04-1 only
    binder = json.loads(ssk_export_binder(scope="04-1", output_dir=str(tmp_path / "binder")))
    assert binder["controls_exported"] == 1
    out = Path(binder["output_dir"])
    assert (out / "04_1.md").exists()
    assert (out / "index.md").exists()
    assert (out / "manifest.json").exists()

    # Coverage reflects 2 controls in DB
    cov = json.loads(ssk_coverage())
    assert cov["total_count"] == 2
    assert isinstance(cov["uncovered"], list)
