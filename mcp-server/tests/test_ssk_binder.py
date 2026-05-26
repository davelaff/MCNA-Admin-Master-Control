import json
from pathlib import Path
import pytest
from docx import Document

from tools.ssk import ssk_import_catalog
from tools.ssk_binder import ssk_coverage, ssk_export_binder


# --------------------------------------------------------------------------- #
# Helpers                                                                       #
# --------------------------------------------------------------------------- #

def _build_control_docx(tmp_path, controls):
    """Build a .docx with one or more controls in MCNA consolidated format."""
    p = tmp_path / "controls.docx"
    doc = Document()
    doc.add_heading("MCNA Cyber Security Standards", level=1)
    for ctrl_id, title in controls:
        gnum = ctrl_id.split("-")[0]
        doc.add_heading(
            f"{ctrl_id}: {title}SSK Group {gnum}: Test StandardsNOF MCNA...", level=1
        )
        for key, val in [("Effective Date", "2026-05-15"), ("Review Date", "2027-05-15"),
                         ("Approver", "CFO / Executive Sponsor")]:
            doc.add_paragraph(key)
            doc.add_paragraph(val)
        doc.add_heading("1. Purpose", level=3)
        doc.add_paragraph(f"{title} overview text.")
        doc.add_heading("4. Standards", level=3)
        doc.add_paragraph("The obligations below apply to MCNA.")
        doc.add_heading("G", level=4)
        doc.add_paragraph(f"{ctrl_id}.1{title} action one.")
        doc.add_heading("6. Review & Compliance", level=3)
        doc.add_heading("Cadence", level=4)
        doc.add_paragraph(f"{title} status text.")
        doc.add_heading("Reviewer", level=4)
        doc.add_paragraph("IT-MIS Director.")
        doc.add_heading("Audit evidence", level=4)
        p2 = doc.add_paragraph(f"{title} evidence.")
        p2.style = doc.styles["List Paragraph"]
        doc.add_heading("7. Risks of Non-compliance", level=3)
        doc.add_paragraph(f"{title} risk text.")
    doc.save(p)
    return p


def _seed_controls(db, tmp_path, controls):
    p = _build_control_docx(tmp_path, controls)
    ssk_import_catalog(
        str(p),
        version="binder-test",
        dry_run=False,
        intermediate_dir=str(tmp_path / "intermediate"),
    )


# --------------------------------------------------------------------------- #
# ssk_coverage tests                                                            #
# --------------------------------------------------------------------------- #

def test_ssk_coverage_returns_expected_keys(db, tmp_path):
    _seed_controls(db, tmp_path, [("08-1", "Identity control"), ("15-3", "PP env control")])
    result = json.loads(ssk_coverage())
    assert "automated_controls" in result
    assert "all_control_ids" in result
    assert "covered_count" in result
    assert "total_count" in result
    assert "uncovered" in result


def test_ssk_coverage_no_controls_imported(db):
    result = json.loads(ssk_coverage())
    assert result["total_count"] == 0
    assert result["covered_count"] == 0
    assert result["automated_controls"] == {}


def test_ssk_coverage_identifies_automated_controls(db, tmp_path):
    # 08-1 is covered by entra (IAM-APP-01, IAM-GUEST-01 etc.) and ca modules
    # 15-3 is covered by pp (PP-ENV-01)
    _seed_controls(db, tmp_path, [("08-1", "Identity control"), ("15-3", "PP env control")])
    result = json.loads(ssk_coverage())
    assert "08-1" in result["automated_controls"]
    assert "15-3" in result["automated_controls"]
    assert result["covered_count"] == 2
    assert result["total_count"] == 2
    assert result["uncovered"] == []


def test_ssk_coverage_uncovered_control_listed(db, tmp_path):
    # 10-1 has no CONTRIBUTES_TO entries in any automated module
    _seed_controls(db, tmp_path, [("10-1", "Uncovered control")])
    result = json.loads(ssk_coverage())
    assert "10-1" in result["uncovered"]
    assert result["covered_count"] == 0


# --------------------------------------------------------------------------- #
# ssk_export_binder — single control                                           #
# --------------------------------------------------------------------------- #

def test_ssk_export_binder_single_creates_output_dir(db, tmp_path):
    _seed_controls(db, tmp_path, [("06-3", "Asset control")])
    out = tmp_path / "binder-out"
    result = json.loads(ssk_export_binder("06-3", output_dir=str(out)))
    assert result["controls_exported"] == 1
    assert Path(result["output_dir"]).exists()


def test_ssk_export_binder_creates_index_md(db, tmp_path):
    _seed_controls(db, tmp_path, [("06-3", "Asset control")])
    out = tmp_path / "binder-out"
    ssk_export_binder("06-3", output_dir=str(out))
    assert (out / "index.md").exists()
    text = (out / "index.md").read_text(encoding="utf-8")
    assert "06-3" in text


def test_ssk_export_binder_creates_manifest_json(db, tmp_path):
    _seed_controls(db, tmp_path, [("06-3", "Asset control")])
    out = tmp_path / "binder-out"
    ssk_export_binder("06-3", output_dir=str(out))
    manifest_path = out / "manifest.json"
    assert manifest_path.exists()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["scope"] == "06-3"
    assert len(manifest["controls"]) == 1
    assert manifest["controls"][0]["control_id"] == "06-3"
    assert "output_dir" in manifest
    assert "generated_at" in manifest


def test_ssk_export_binder_creates_control_md_file(db, tmp_path):
    _seed_controls(db, tmp_path, [("06-3", "Asset control")])
    out = tmp_path / "binder-out"
    ssk_export_binder("06-3", output_dir=str(out))
    ctrl_file = out / "06_3.md"
    assert ctrl_file.exists()
    text = ctrl_file.read_text(encoding="utf-8")
    assert "06-3" in text
    assert "Asset control" in text


def test_ssk_export_binder_md_contains_required_sections(db, tmp_path):
    _seed_controls(db, tmp_path, [("06-3", "Asset control")])
    out = tmp_path / "binder-out"
    ssk_export_binder("06-3", output_dir=str(out))
    text = (out / "06_3.md").read_text(encoding="utf-8")
    for section in [
        "MCNA Position",
        "Control Statement",
        "Standards Clauses",
        "Audit Evidence Requirements",
        "Evidence",
        "Review History",
        "Active Findings",
        "Insufficient Measures Risks",
    ]:
        assert section in text, f"Missing required section: {section}"


def test_ssk_export_binder_md_contains_overview_text(db, tmp_path):
    _seed_controls(db, tmp_path, [("06-3", "Asset control")])
    out = tmp_path / "binder-out"
    ssk_export_binder("06-3", output_dir=str(out))
    text = (out / "06_3.md").read_text(encoding="utf-8")
    assert "Asset control overview text." in text


# --------------------------------------------------------------------------- #
# ssk_export_binder — all scope                                                #
# --------------------------------------------------------------------------- #

def test_ssk_export_binder_all_scope(db, tmp_path):
    _seed_controls(db, tmp_path, [("04-1", "HR control"), ("06-3", "Asset control")])
    out = tmp_path / "binder-all"
    result = json.loads(ssk_export_binder("all", output_dir=str(out)))
    assert result["controls_exported"] == 2
    assert (out / "04_1.md").exists()
    assert (out / "06_3.md").exists()


def test_ssk_export_binder_all_index_lists_all_controls(db, tmp_path):
    _seed_controls(db, tmp_path, [("04-1", "HR control"), ("06-3", "Asset control")])
    out = tmp_path / "binder-all"
    ssk_export_binder("all", output_dir=str(out))
    text = (out / "index.md").read_text(encoding="utf-8")
    assert "04-1" in text
    assert "06-3" in text


def test_ssk_export_binder_reused_output_dir_merges_index_and_manifest(db, tmp_path):
    _seed_controls(db, tmp_path, [("06-3", "Asset control"), ("08-1", "Identity control")])
    out = tmp_path / "binder-merge"

    first = json.loads(ssk_export_binder("06-3", output_dir=str(out)))
    second = json.loads(ssk_export_binder("08-1", output_dir=str(out)))

    assert first["controls_exported"] == 1
    assert second["controls_exported"] == 1
    assert (out / "06_3.md").exists()
    assert (out / "08_1.md").exists()

    index_text = (out / "index.md").read_text(encoding="utf-8")
    assert "06-3" in index_text
    assert "08-1" in index_text

    manifest = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["scope"] == "custom-multi-control"
    assert [entry["control_id"] for entry in manifest["controls"]] == ["06-3", "08-1"]


# --------------------------------------------------------------------------- #
# ssk_export_binder — error cases                                               #
# --------------------------------------------------------------------------- #

def test_ssk_export_binder_unknown_control_returns_error(db, tmp_path):
    out = tmp_path / "binder-err"
    result = json.loads(ssk_export_binder("99-9", output_dir=str(out)))
    assert "error" in result


def test_ssk_export_binder_returns_json_ok_payload(db, tmp_path):
    _seed_controls(db, tmp_path, [("06-3", "Asset control")])
    out = tmp_path / "binder-check"
    result = json.loads(ssk_export_binder("06-3", output_dir=str(out)))
    assert "output_dir" in result
    assert "controls_exported" in result
    assert "files" in result
    assert isinstance(result["files"], list)
    assert len(result["files"]) == 1
